"""Экспорт книги: .md, .txt, .docx."""

import io
import re
from pathlib import Path

import db

EXPORT_DIR = Path(__file__).parent / "exports"
EXPORT_DIR.mkdir(exist_ok=True)

# Подписи на языке КНИГИ (project.language), не языка интерфейса.
_LABELS = {
    "ru": {
        "untitled":      "Без названия",
        "genre":         "Жанр",
        "setting":       "Место и время",
        "tone":          "Тональность",
        "chapters":      "Глав",
        "words":         "слов",
        "chars":         "знаков",
        "chapter":       "Глава {n}.",
        "not_generated": "(глава ещё не сгенерирована)",
    },
    "en": {
        "untitled":      "Untitled",
        "genre":         "Genre",
        "setting":       "Place and time",
        "tone":          "Tone",
        "chapters":      "Chapters",
        "words":         "words",
        "chars":         "chars",
        "chapter":       "Chapter {n}.",
        "not_generated": "(chapter not generated yet)",
    },
}


def _L(key: str) -> str:
    lang = (db.get_project() or {}).get("language", "ru")
    return _LABELS.get(lang, _LABELS["ru"])[key]


def _safe_name(s: str) -> str:
    s = re.sub(r"[^\w\s\-а-яА-ЯёЁ]", "", s or "")
    s = re.sub(r"\s+", "_", s.strip())
    return s[:60] or "novel"


def _meta() -> dict:
    project = db.get_project() or {}
    return {
        "title":   project.get("theme") or _L("untitled"),
        "genre":   project.get("genre") or "",
        "setting": project.get("setting") or "",
        "tone":    project.get("tone") or "",
    }


def _chapters() -> list[dict]:
    with db.conn() as c:
        rows = c.execute(
            "SELECT num, title, pov_char, text, word_count, status "
            "FROM chapters ORDER BY num"
        ).fetchall()
    return [dict(r) for r in rows]


def _stats(chapters: list[dict]) -> dict:
    total_words = sum(ch.get("word_count") or 0 for ch in chapters)
    total_chars = sum(len(ch.get("text") or "") for ch in chapters)
    return {
        "chapters": len(chapters),
        "words":    total_words,
        "chars":    total_chars,
        "done":     sum(1 for ch in chapters if ch.get("status") == "accepted"),
        "draft":    sum(1 for ch in chapters if ch.get("status") == "draft"),
    }


# ────────────────────────────────────────────────────────────────
# Markdown
# ────────────────────────────────────────────────────────────────

def to_markdown() -> str:
    meta = _meta()
    chapters = _chapters()
    s = _stats(chapters)

    parts = [
        f"# {meta['title']}\n",
        f"*{_L('genre')}:* {meta['genre']} · *{_L('setting')}:* {meta['setting']}"
        + (f" · *{_L('tone')}:* {meta['tone']}" if meta["tone"] else ""),
        "",
        f"*{_L('chapters')}:* {s['chapters']} · *{_L('words')}:* {s['words']}"
        f" · *{_L('chars')}:* {s['chars']}",
        "",
        "---",
        "",
    ]

    for ch in chapters:
        parts.append(f"## {_L('chapter').format(n=ch['num'])} {ch.get('title') or ''}")
        parts.append("")
        body = (ch.get("text") or "").strip()
        if body:
            parts.append(body)
        else:
            parts.append(f"*{_L('not_generated')}*")
        parts.append("")
        parts.append("---")
        parts.append("")

    return "\n".join(parts)


# ────────────────────────────────────────────────────────────────
# Plain text
# ────────────────────────────────────────────────────────────────

def to_txt() -> str:
    meta = _meta()
    chapters = _chapters()
    lines = [
        meta["title"],
        "=" * len(meta["title"]),
        "",
        f"{_L('genre')}: {meta['genre']}",
        f"{_L('setting')}: {meta['setting']}",
    ]
    if meta["tone"]:
        lines.append(f"{_L('tone')}: {meta['tone']}")
    lines.append("")
    lines.append("")
    for ch in chapters:
        lines.append(f"{_L('chapter').format(n=ch['num'])} {ch.get('title') or ''}")
        lines.append("-" * 40)
        lines.append((ch.get("text") or "").strip() or _L("not_generated"))
        lines.append("")
        lines.append("")
    return "\n".join(lines)


# ────────────────────────────────────────────────────────────────
# DOCX
# ────────────────────────────────────────────────────────────────

def to_docx_bytes() -> bytes:
    from docx import Document
    from docx.shared import Pt
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    meta = _meta()
    chapters = _chapters()

    doc = Document()

    # Стиль по умолчанию
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(12)

    # Титул
    h = doc.add_heading(meta["title"], level=0)
    h.alignment = WD_ALIGN_PARAGRAPH.CENTER

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = sub.add_run(
        f"{meta['genre']} · {meta['setting']}"
        + (f" · {meta['tone']}" if meta["tone"] else "")
    )
    run.italic = True

    doc.add_page_break()

    for ch in chapters:
        doc.add_heading(f"{_L('chapter').format(n=ch['num'])} {ch.get('title') or ''}",
                        level=1)
        body = (ch.get("text") or "").strip() or _L("not_generated")
        for para in body.split("\n\n"):
            para = para.strip()
            if not para:
                continue
            p = doc.add_paragraph(para)
            p.paragraph_format.first_line_indent = Pt(24)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ────────────────────────────────────────────────────────────────
# Сохранение на диск
# ────────────────────────────────────────────────────────────────

def save_all() -> dict:
    """Сохранить все форматы в exports/ и вернуть пути."""
    name = _safe_name(_meta()["title"])
    paths = {}

    p_md = EXPORT_DIR / f"{name}.md"
    p_md.write_text(to_markdown(), encoding="utf-8")
    paths["md"] = str(p_md)

    p_txt = EXPORT_DIR / f"{name}.txt"
    p_txt.write_text(to_txt(), encoding="utf-8")
    paths["txt"] = str(p_txt)

    p_docx = EXPORT_DIR / f"{name}.docx"
    p_docx.write_bytes(to_docx_bytes())
    paths["docx"] = str(p_docx)

    return paths